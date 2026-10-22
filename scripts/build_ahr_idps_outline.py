from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "Đề cương báo cáo bài tập lớn AHR-IDPS.docx"


def set_run_font(run, size: float = 13, bold: bool | None = None, italic: bool = False):
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold
    run.italic = italic


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="6"):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def mark_header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def add_field(paragraph, instruction: str):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_char3 = OxmlElement("w:fldChar")
    fld_char3.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2, text, fld_char3])


def add_page_number(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Trang ")
    set_run_font(r, 10)
    add_field(p, "PAGE")


def add_page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def add_body(doc, text: str, *, italic=False):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.first_line_indent = Cm(0.8)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    set_run_font(r, 13, italic=italic)
    return p


def add_bullet(doc, text: str):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.1
    r = p.add_run(text)
    set_run_font(r, 13)
    return p


def add_heading(doc, text: str, level: int = 1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(12 if level == 1 else 7)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    set_run_font(r, 15 if level == 1 else 13, bold=True)
    return p


def add_caption(doc, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    set_run_font(r, 11, italic=True)
    return p


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table)
    hdr = table.rows[0].cells
    mark_header_row(table.rows[0])
    for i, h in enumerate(headers):
        hdr[i].text = ""
        set_cell_shading(hdr[i], "404040")
        hdr[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(hdr[i])
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        set_run_font(r, 11, bold=True)
        r.font.color.rgb = RGBColor(255, 255, 255)
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = ""
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cells[i])
            if row_index % 2 == 1:
                set_cell_shading(cells[i], "F2F2F2")
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(str(value))
            set_run_font(r, 10.5)
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(13)
    for name, size in (("Title", 18), ("Heading 1", 15), ("Heading 2", 13), ("Heading 3", 13)):
        style = styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
    if "Subtitle" in styles:
        styles["Subtitle"].font.name = "Times New Roman"
        styles["Subtitle"].font.size = Pt(14)
        styles["Subtitle"].font.color.rgb = RGBColor(0, 0, 0)


def main():
    doc = Document()
    configure_styles(doc)
    for section in doc.sections:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(2.0)
        add_page_number(section)

    # Cover
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(25)
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("ĐỀ CƯƠNG BÁO CÁO BÀI TẬP LỚN")
    set_run_font(r, 18, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("XÂY DỰNG HỆ THỐNG AHR-IDPS")
    set_run_font(r, 16, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(22)
    r = p.add_run("Hybrid Intrusion Detection and Response System kết hợp luật tự sinh và mô hình AI thích nghi cho giám sát an ninh và ứng phó xâm nhập trên ứng dụng web DVWA")
    set_run_font(r, 13, bold=True)
    for line in (
        "Môn học: An toàn ứng dụng Web",
        "Đối tượng thực nghiệm: Ứng dụng DVWA trong môi trường Docker cục bộ",
        "Nhóm thực hiện: ........................................................",
        "Giảng viên hướng dẫn: ....................................................",
        "Năm học: 2025 - 2026",
    ):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(8)
        r = p.add_run(line)
        set_run_font(r, 13)
    add_page_break(doc)

    # TOC/front matter
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("MỤC LỤC")
    set_run_font(r, 16, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("Mục lục được tạo từ các tiêu đề Heading 1 đến Heading 3. Trong Microsoft Word, chọn Update Field để cập nhật số trang sau khi bổ sung nội dung.")
    set_run_font(r, 11, italic=True)
    p = doc.add_paragraph()
    add_field(p, 'TOC \\o "1-3" \\h \\z \\u')
    add_page_break(doc)

    add_heading(doc, "DANH MỤC HÌNH ẢNH", 1)
    add_body(doc, "Danh mục dưới đây là danh sách hình cần xây dựng hoặc cập nhật trong báo cáo cuối. Các hình nên được đánh số thống nhất và có nguồn dữ liệu rõ ràng.")
    add_table(doc, ["Mã hình", "Tên hình dự kiến", "Nguồn / nơi lấy"], [
        ("Hình 3.1", "Kiến trúc tổng thể AHR-IDPS", "docs/visuals/ahr-idps-system-overview.png"),
        ("Hình 3.2", "Luồng sự kiện từ DVWA đến incident", "Sơ đồ triển khai và pipeline"),
        ("Hình 3.3", "Sơ đồ attack graph và các node quan trọng", "dashboard/src/components/AttackGraph.tsx"),
        ("Hình 5.1", "Luồng dữ liệu huấn luyện và suy luận AI", "docs/visuals/ahr-idps-ai-flow.png"),
        ("Hình 6.1", "Giao diện dashboard giám sát", "Ảnh chụp dashboard thực tế"),
        ("Hình 6.2", "Giao diện quản trị Wazuh manager", "Ảnh chụp API/console loopback"),
        ("Hình 7.1", "So sánh kết quả rule, AI và hybrid", "Biểu đồ thực nghiệm"),
    ], widths=[2.3, 9.2, 5.0])
    add_heading(doc, "DANH MỤC BẢNG", 1)
    add_table(doc, ["Mã bảng", "Tên bảng dự kiến", "Nội dung cần điền"], [
        ("Bảng 2.1", "Các chỉ số đánh giá hệ thống", "Precision, recall, F1, FPR, latency, throughput"),
        ("Bảng 3.1", "Vai trò các node trong mô hình", "Input, xử lý, output, trạng thái lỗi"),
        ("Bảng 5.1", "Cấu trúc dataset AHR-IDPS", "Scenario, label, event type, timestamp"),
        ("Bảng 6.1", "Ma trận kiểm thử giao diện và cổng", "URL, expected status, kết quả thực tế"),
        ("Bảng 6.2", "Ba kịch bản pentest được ủy quyền", "Mục tiêu, lệnh, tín hiệu mong đợi, bằng chứng"),
        ("Bảng 7.1", "Kết quả phát hiện theo từng sensor", "Suricata, Falco, Wazuh, AI, fusion"),
    ], widths=[2.3, 9.2, 5.0])
    add_page_break(doc)

    add_heading(doc, "LỜI MỞ ĐẦU", 1)
    add_body(doc, "Đề tài này xây dựng một hệ thống Hybrid Intrusion Detection and Response System, gọi tắt là AHR-IDPS, để giám sát và ứng phó xâm nhập trên ứng dụng web DVWA trong môi trường kiểm thử cục bộ. Báo cáo cần làm rõ từ cơ sở lý thuyết, kiến trúc node, pipeline dữ liệu, cách sinh luật, huấn luyện AI, kiểm thử pentest đến đánh giá định lượng.")
    add_body(doc, "Đề cương này được dùng làm khung viết báo cáo cuối. Những vị trí ghi kết quả, ảnh chụp, số liệu hoặc phiên bản cần được thay bằng bằng chứng thu được khi chạy hệ thống; không điền số liệu giả định nếu chưa thực nghiệm.")

    add_heading(doc, "CHƯƠNG 1. GIỚI THIỆU ĐỀ TÀI", 1)
    for h in [
        "1.1. Bối cảnh và động lực nghiên cứu",
        "1.2. Vấn đề cần giải quyết",
        "1.3. Mục tiêu tổng quát",
        "1.4. Mục tiêu cụ thể",
        "1.5. Câu hỏi nghiên cứu",
        "1.6. Đối tượng và phạm vi nghiên cứu",
        "1.7. Phương pháp nghiên cứu",
        "1.8. Đóng góp của đề tài",
        "1.9. Cấu trúc báo cáo",
    ]:
        add_heading(doc, h, 2)
        add_body(doc, "Nêu rõ nội dung cần chứng minh, nguồn dữ liệu sử dụng, điều kiện thực nghiệm và tiêu chí hoàn thành của mục này.", italic=True)

    add_heading(doc, "CHƯƠNG 2. CƠ SỞ LÝ THUYẾT VÀ CÔNG NGHỆ", 1)
    sections = {
        "2.1. An ninh ứng dụng Web và OWASP Top 10": ["Khái niệm lỗ hổng web", "SQL Injection, XSS, Command Injection, File Inclusion và Authentication Abuse", "Nguyên tắc kiểm thử có ủy quyền"],
        "2.2. Ứng dụng DVWA và môi trường lab": ["Mục đích của DVWA", "Các mức bảo mật và giới hạn lab", "Luồng đăng nhập và session"],
        "2.3. IDS, IPS và IDRS": ["Phát hiện, phân loại, phản ứng", "So sánh signature-based, anomaly-based và hybrid"],
        "2.4. Phát hiện dựa trên luật": ["Luật Suricata, Falco, Wazuh", "Luật tự sinh và kiểm duyệt trước khi áp dụng", "Giảm false positive"],
        "2.5. Phát hiện bất thường bằng AI": ["Đặc trưng mạng và log", "XGBoost, Isolation Forest/MLP, embedding", "Thích nghi theo cửa sổ dữ liệu"],
        "2.6. Tương quan sự kiện và attack graph": ["Sessionization", "Correlation và risk scoring", "GraphSAGE/window graph"],
        "2.7. SIEM, telemetry và evidence": ["Redis Streams, PostgreSQL", "Evidence bundle và SHA-256", "Auditability"],
        "2.8. Chỉ số đánh giá": ["Precision, recall, F1, FPR", "ROC-AUC/PR-AUC", "Latency, throughput và detection coverage"],
    }
    for h, subs in sections.items():
        add_heading(doc, h, 2)
        for s in subs:
            add_heading(doc, s, 3)
            add_body(doc, "Trình bày khái niệm, công thức hoặc sơ đồ liên quan và liên hệ trực tiếp với AHR-IDPS.", italic=True)

    add_heading(doc, "CHƯƠNG 3. PHÂN TÍCH YÊU CẦU VÀ THIẾT KẾ HỆ THỐNG", 1)
    for h, body in [
        ("3.1. Yêu cầu chức năng", "Thu nhận sự kiện; chuẩn hóa; phát hiện bằng luật và AI; tương quan; sinh incident; hiển thị dashboard; quản trị Wazuh; tạo evidence; phản ứng ở chế độ alert_only."),
        ("3.2. Yêu cầu phi chức năng", "Tính ổn định khi Docker restart, khả năng truy vết, bảo vệ loopback, khả năng mở rộng, thời gian phản hồi và tái lập thí nghiệm."),
        ("3.3. Mô hình đe dọa và giả định", "Tác nhân là máy khách trong lab; DVWA là tài sản cần bảo vệ; không pentest ra Internet; marker được dùng để truy vết."),
        ("3.4. Kiến trúc tổng thể", "Chèn Hình 3.1 và mô tả ranh giới DVWA, sensors, pipeline, AI, fusion, response, dashboard."),
        ("3.5. Chức năng của từng node", "Lập bảng theo input, xử lý, output, cổng, health check và lỗi thường gặp cho DVWA/Nginx/MariaDB, Suricata, Falco, Wazuh, Collector, Normalizer, Redis, PostgreSQL, Backend, Dashboard, AI, Fusion, Incident, Rule Generator, Response và Evidence."),
        ("3.6. Luồng dữ liệu và luồng quyết định", "Mô tả đường đi từ request DVWA đến alert, risk score, incident, rule đề xuất và response."),
        ("3.7. Thiết kế schema sự kiện", "Trình bày các trường event_id, timestamp, source, scenario_id, session_id, label, features, severity, model_version và evidence_ref."),
        ("3.8. Thiết kế phân quyền và ranh giới quản trị", "Nêu các cổng loopback, API key, quyền read-only, dashboard role và quản trị Wazuh manager."),
    ]:
        add_heading(doc, h, 2)
        add_body(doc, body)
    add_table(doc, ["Node", "Chức năng chính", "Đầu vào", "Đầu ra / kiểm tra"], [
        ("DVWA + Nginx + MariaDB", "Tạo ứng dụng và traffic lab", "HTTP request, login, database", "Access log, response, health endpoint"),
        ("Suricata", "Phân tích mạng và luật", "Packet/HTTP", "EVE JSON, alert, flow"),
        ("Falco", "Giám sát runtime/eBPF", "Syscall/container event", "Runtime alert"),
        ("Wazuh manager", "Quản lý alert và rule", "Log/JSON alert", "Alert, status, manager API"),
        ("Collector + Normalizer", "Thu nhận và chuẩn hóa", "Sensor outputs", "Canonical event"),
        ("Redis + PostgreSQL", "Stream và lưu trữ", "Canonical event", "Queue, query, replay"),
        ("AI + Fusion + Incident", "Scoring và tương quan", "Features/window graph", "Risk, incident, explanation"),
        ("Dashboard + Response", "Quan sát và phản ứng", "Incident/evidence", "UI, alert_only action, audit"),
    ], widths=[3.0, 5.0, 4.0, 4.5])

    add_heading(doc, "CHƯƠNG 4. TRIỂN KHAI HỆ THỐNG", 1)
    for h in [
        "4.1. Cấu hình phần cứng, hệ điều hành và Docker",
        "4.2. Triển khai DVWA, Nginx và MariaDB",
        "4.3. Triển khai Suricata và kiểm tra EVE JSON",
        "4.4. Triển khai Falco/eBPF và kiểm tra runtime event",
        "4.5. Triển khai Wazuh manager và API loopback",
        "4.6. Collector, Normalizer, Redis Streams và PostgreSQL",
        "4.7. FastAPI backend, WebSocket và React dashboard",
        "4.8. Các dịch vụ AI, Fusion Engine và Incident Engine",
        "4.9. Rule Generator và quy trình phê duyệt luật",
        "4.10. Response Engine ở chế độ alert_only",
        "4.11. Evidence bundle, manifest và SHA-256",
        "4.12. Quy trình khởi động, health gate và xử lý lỗi Docker",
    ]:
        add_heading(doc, h, 2)
        add_body(doc, "Ghi phiên bản image, biến môi trường, cổng bind, health check, lệnh khởi động, log lỗi và cách khắc phục đã xác nhận.", italic=True)

    add_heading(doc, "CHƯƠNG 5. DATASET VÀ HUẤN LUYỆN MÔ HÌNH AI", 1)
    for h, body in [
        ("5.1. Mục tiêu dataset", "Dataset phải phản ánh traffic DVWA thật, giữ được session/scenario và có thể truy nguyên về evidence."),
        ("5.2. Thiết kế scenario và nhãn", "Baseline, web_injection và auth_abuse; nhãn benign/suspicious/attack; marker và metadata."),
        ("5.3. Thu thập dữ liệu", "Dùng authorized_monitoring_scenarios.py và collect_dvwa_training_dataset.ps1; lưu provenance, version và timestamp."),
        ("5.4. Làm sạch và chuẩn hóa", "Loại bỏ duplicate, xử lý missing, đồng bộ timestamp, chuẩn hóa categorical và giữ event ordering."),
        ("5.5. Chia train/validation/test", "Ưu tiên split theo scenario_id để tránh leakage; đề xuất 60/15/15/10 cho train/validation/test/holdout."),
        ("5.6. Dataset bổ sung", "Có thể dùng CIC-IDS2017, tập Web Attack Brute Force/Web Attack XSS/Web Attack SQL Injection để bổ sung; phải ghi rõ khác biệt schema và không trộn mù với dữ liệu DVWA."),
        ("5.7. Huấn luyện Network AI", "Đặc trưng flow/HTTP, XGBoost, cân bằng lớp, lưu model artifact và metrics."),
        ("5.8. Huấn luyện Log AI", "Embedding hoặc vector log, Isolation Forest/MLP, ngưỡng anomaly và kiểm tra drift."),
        ("5.9. Huấn luyện Graph AI", "Window graph, node/edge feature, GraphSAGE hoặc baseline graph heuristic."),
        ("5.10. Fusion và giải thích", "Kết hợp điểm sensor, rule và AI; lưu lý do, model_version, threshold và confidence."),
        ("5.11. Đăng ký model và tái lập", "Dùng register_existing_models.sql, metadata JSON, checksum và quy trình rollback."),
    ]:
        add_heading(doc, h, 2)
        add_body(doc, body)
    add_table(doc, ["Trường", "Ví dụ", "Ý nghĩa"], [
        ("scenario_id", "web_injection_001", "Nhóm thí nghiệm và khóa split"),
        ("label", "benign / attack", "Nhãn dùng huấn luyện và đánh giá"),
        ("event_type", "nginx_access / suricata_alert", "Nguồn và loại telemetry"),
        ("session_id", "lab-session-01", "Liên kết chuỗi request"),
        ("model_version", "network-xgb-v1", "Tái lập model và rollback"),
    ], widths=[4.0, 5.0, 8.5])

    add_heading(doc, "CHƯƠNG 6. THIẾT KẾ KIỂM THỬ VÀ QUY TRÌNH GIÁM SÁT", 1)
    add_body(doc, "Chương này phải mô tả bài kiểm thử có ủy quyền, chỉ chạy trên DVWA local/loopback, có giới hạn tốc độ và lưu bằng chứng. Mỗi kịch bản cần ghi thời gian bắt đầu, scenario_id, tài khoản lab, lệnh, expected result, actual result và đường dẫn evidence.")
    for h, body in [
        ("6.1. Kiểm thử môi trường và cổng dịch vụ", "Kiểm tra Docker daemon, container health, HTTP status, API, WebSocket, dashboard, Wazuh manager console và các cổng loopback đã công bố."),
        ("6.2. Kịch bản 1 - Baseline hợp lệ", "Đăng nhập DVWA, mở trang chủ/instructions và thực hiện request bình thường. Kỳ vọng: telemetry đầy đủ, không tạo incident nghiêm trọng, đo false positive."),
        ("6.3. Kịch bản 2 - Web injection có kiểm soát", "Sau login, thực hiện tín hiệu SQL injection, reflected XSS, command injection và file inclusion với marker, tốc độ thấp. Kỳ vọng: rule/sensor/AI ghi nhận và fusion tạo incident."),
        ("6.4. Kịch bản 3 - Authentication abuse", "Thực hiện số lần login sai có giới hạn và truy cập path nhạy cảm trong lab. Kỳ vọng: phát hiện pattern lặp, sessionization và risk tăng theo chuỗi sự kiện."),
        ("6.5. Kiểm thử từng sensor", "Đối chiếu Nginx, Suricata, Falco và Wazuh với cùng một scenario_id; ghi nhận sensor nào có/không có event."),
        ("6.6. Kiểm thử dashboard và API", "Kiểm tra view Giám sát, Điều tra, Quản trị; API incidents/events/models; WebSocket reconnect; trạng thái thiếu dữ liệu phải hiển thị rõ."),
        ("6.7. Quy trình giám sát và triage", "Alert -> kiểm tra timeline -> xem attack graph -> đối chiếu evidence -> xác nhận/đóng incident -> ghi nhận phản hồi."),
        ("6.8. Thu thập evidence", "Lưu raw event, normalized event, alert, incident, model decision, rule decision và manifest SHA-256."),
    ]:
        add_heading(doc, h, 2)
        add_body(doc, body)
    add_table(doc, ["Kịch bản", "Tác vụ", "Kết quả mong đợi", "Bằng chứng"], [
        ("baseline", "Login và truy cập hợp lệ", "Không false positive nghiêm trọng; có access/flow event", "scenario marker, event count, dashboard"),
        ("web_injection", "SQLi/XSS/command/file inclusion nhẹ", "Có alert/risk/incident; giữ nguyên hệ thống", "Suricata/Wazuh/Falco/AI/evidence"),
        ("auth_abuse", "Login sai giới hạn và path nhạy cảm", "Có correlation theo session và risk tăng", "auth event, timeline, incident"),
    ], widths=[3.0, 5.2, 6.1, 3.2])

    add_heading(doc, "CHƯƠNG 7. KẾT QUẢ VÀ ĐÁNH GIÁ", 1)
    for h in [
        "7.1. Kết quả triển khai và độ ổn định",
        "7.2. Mức độ bao phủ của các sensor",
        "7.3. Kết quả kiểm tra cổng và giao diện",
        "7.4. Precision, recall, F1, FPR và PR-AUC",
        "7.5. Latency, throughput và tài nguyên Docker",
        "7.6. So sánh rule-only, AI-only và hybrid",
        "7.7. Thử nghiệm unknown attack và holdout scenario",
        "7.8. Kết quả incident, evidence và response alert_only",
        "7.9. Phân tích false positive và cách điều chỉnh ngưỡng",
        "7.10. Hạn chế và nguy cơ ảnh hưởng kết quả",
    ]:
        add_heading(doc, h, 2)
        add_body(doc, "Điền số liệu thực nghiệm, bảng kết quả, hình dashboard và nhận xét có đối chiếu với tiêu chí ở Chương 1.", italic=True)
    add_table(doc, ["Thử nghiệm", "Cấu hình", "Chỉ số cần báo cáo"], [
        ("E1 - Rule only", "Suricata/Falco/Wazuh", "TP, FP, FN, TN, latency"),
        ("E2 - Network AI", "XGBoost", "F1, PR-AUC, inference time"),
        ("E3 - Log AI", "Embedding + anomaly", "F1, FPR, drift"),
        ("E4 - Graph AI", "Window graph", "Incident coverage, lead time"),
        ("E5 - Fusion", "Rule + 3 AI", "F1, latency, explanation"),
        ("E6-E9", "Ablation/holdout", "Độ nhạy và khả năng khái quát"),
    ], widths=[4.0, 5.5, 8.0])

    add_heading(doc, "CHƯƠNG 8. KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN", 1)
    for h, body in [
        ("8.1. Kết luận", "Tóm tắt hệ thống đã xây dựng, luồng giám sát, kết quả kiểm thử và mức độ đáp ứng mục tiêu."),
        ("8.2. Đóng góp thực tiễn", "Nêu giá trị của pipeline thu thập, hybrid decision, dashboard, evidence và runbook pentest có ủy quyền."),
        ("8.3. Hạn chế", "Nêu giới hạn về dữ liệu, Wazuh manager-only, alert_only, môi trường Docker cục bộ và kích thước dataset."),
        ("8.4. Hướng phát triển", "Mở rộng Wazuh indexer/dashboard, drift monitoring, active response có phê duyệt, dataset lớn hơn và triển khai phân tán."),
    ]:
        add_heading(doc, h, 2)
        add_body(doc, body)

    add_heading(doc, "TÀI LIỆU THAM KHẢO", 1)
    for ref in [
        "OWASP Foundation, OWASP Top 10 và Web Security Testing Guide.",
        "Damn Vulnerable Web Application, tài liệu và mã nguồn DVWA.",
        "Suricata Documentation, EVE JSON và rule language.",
        "Falco Documentation, rules và runtime security.",
        "Wazuh Documentation, manager, rules, API và alert schema.",
        "CIC-IDS2017, Canadian Institute for Cybersecurity, tập dữ liệu bổ sung tham khảo.",
        "Tài liệu nội bộ của dự án AHR-IDPS: README, runbook, dataset training guide và source code.",
    ]:
        add_bullet(doc, ref)

    add_heading(doc, "PHỤ LỤC A. CẤU HÌNH VÀ LỆNH VẬN HÀNH", 1)
    for h in ["A.1. Compose và biến môi trường", "A.2. Health check và kiểm tra cổng", "A.3. API endpoint và WebSocket", "A.4. Cấu hình Wazuh manager loopback"]:
        add_heading(doc, h, 2)
        add_body(doc, "Chèn đoạn cấu hình đã dùng, phiên bản, lệnh chạy và kết quả kiểm tra tương ứng.", italic=True)

    add_heading(doc, "PHỤ LỤC B. BA KỊCH BẢN PENTEST ĐƯỢC ỦY QUYỀN", 1)
    add_body(doc, "Chỉ chạy ba kịch bản trên DVWA local theo scripts/authorized_monitoring_scenarios.py. Không thực hiện DoS, persistence, reverse shell hoặc quét ra ngoài phạm vi lab.")
    for h, body in [
        ("B.1. Baseline", "Mục tiêu: tạo traffic hợp lệ và đo false positive. Dữ liệu cần lưu: login marker, access event, flow event, dashboard snapshot."),
        ("B.2. Web injection", "Mục tiêu: xác nhận rule, AI, correlation và evidence đối với tín hiệu SQLi/XSS/command/file inclusion có kiểm soát."),
        ("B.3. Authentication abuse", "Mục tiêu: xác nhận phát hiện login sai lặp lại, sessionization, risk scoring và incident timeline."),
    ]:
        add_heading(doc, h, 2)
        add_body(doc, body)

    add_heading(doc, "PHỤ LỤC C. CHECKLIST HOÀN THIỆN BÁO CÁO", 1)
    for item in [
        "Đã thay mọi đoạn hướng dẫn bằng kết quả thực nghiệm có timestamp và phiên bản.",
        "Đã chèn sơ đồ kiến trúc, luồng AI, attack graph và ảnh dashboard thực tế.",
        "Đã kiểm tra từng node, cổng, API, WebSocket và Wazuh manager.",
        "Đã chạy đủ baseline, web_injection và auth_abuse trong phạm vi DVWA local.",
        "Đã lưu raw event, normalized event, incident và evidence manifest.",
        "Đã tính chỉ số rule-only, AI-only, hybrid và phân tích false positive.",
        "Đã cập nhật mục lục, danh mục hình/bảng, tài liệu tham khảo và phụ lục.",
    ]:
        add_bullet(doc, "☐ " + item)

    # Ask Word to refresh the TOC when opened.
    settings = doc.settings.element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    doc.save(OUTPUT)
    print(str(OUTPUT).encode("unicode_escape").decode("ascii"))


if __name__ == "__main__":
    main()
