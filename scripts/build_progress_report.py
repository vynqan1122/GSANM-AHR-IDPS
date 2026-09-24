from __future__ import annotations

from pathlib import Path
from shutil import copyfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Pt
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = Path(r"D:\GSANM\Báo cáo tiến độ Bài tập lớn.docx")
OUTPUT = ROOT / "Báo cáo tiến độ AHR-IDPS.docx"


def set_font(run, name: str = "Times New Roman", size: int = 13, bold: bool | None = None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def ensure_styles(doc: Document) -> None:
    if "Title" not in [style.name for style in doc.styles]:
        style = doc.styles.add_style("Title", WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "Times New Roman"
        style.font.size = Pt(16)
        style.font.bold = True
    for name in ("Heading 1", "Heading 2"):
        if name not in [style.name for style in doc.styles]:
            style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            style.font.name = "Times New Roman"
            style.font.size = Pt(14 if name == "Heading 1" else 13)
            style.font.bold = True


def replace_paragraph(paragraph, text: str, *, bold: bool = False, size: int = 13, align=None):
    for run in paragraph.runs:
        run.text = ""
    run = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
    run.text = text
    set_font(run, size=size, bold=bold)
    if align is not None:
        paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.15


def add_body(doc: Document, text: str, *, bold: bool = False, italic: bool = False):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.15
    run = paragraph.add_run(text)
    set_font(run, size=13, bold=bold)
    run.italic = italic
    return paragraph


def add_heading(doc: Document, text: str, level: int = 1):
    paragraph = doc.add_paragraph()
    paragraph.style = "Heading 1" if level == 1 else "Heading 2"
    paragraph.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(text)
    set_font(run, size=14 if level == 1 else 13, bold=True)
    return paragraph


def main() -> None:
    if not TEMPLATE.exists():
        raise FileNotFoundError(TEMPLATE)
    copyfile(TEMPLATE, OUTPUT)
    doc = Document(OUTPUT)
    ensure_styles(doc)

    for section in doc.sections:
        section.top_margin = Pt(57)
        section.bottom_margin = Pt(57)
        section.left_margin = Pt(85)
        section.right_margin = Pt(57)

    title = doc.paragraphs[0]
    title.style = "Title"
    replace_paragraph(title, "BÁO CÁO TIẾN ĐỘ THỰC HIỆN ĐỀ TÀI", bold=True, size=16, align=WD_ALIGN_PARAGRAPH.CENTER)
    replace_paragraph(doc.paragraphs[1], "Tên đề tài: Xây dựng hệ thống Hybrid Intrusion Detection and Response System kết hợp luật tự sinh và mô hình AI thích nghi cho giám sát an ninh và ứng phó xâm nhập trên ứng dụng web DVWA", bold=True)
    replace_paragraph(doc.paragraphs[2], "Nhóm sinh viên thực hiện: Nhóm AHR-IDPS")
    replace_paragraph(doc.paragraphs[3], "NỘI DUNG BÁO CÁO", bold=True, size=14, align=WD_ALIGN_PARAGRAPH.CENTER)
    for index in (4, 7, 11, 14):
        doc.paragraphs[index].style = "Heading 1"
    replace_paragraph(doc.paragraphs[4], "1. Mục tiêu và phạm vi nghiên cứu")
    replace_paragraph(doc.paragraphs[5], "Mục tiêu: Xây dựng một hệ thống phát hiện và ứng phó xâm nhập lai cho ứng dụng web DVWA, kết hợp các luật phát hiện với các mô hình AI về mạng, log và đồ thị. Hệ thống phải thu nhận được sự kiện thật, chuẩn hóa dữ liệu, tương quan thành incident và hỗ trợ người vận hành giám sát qua dashboard.")
    replace_paragraph(doc.paragraphs[6], "Phạm vi: Môi trường kiểm thử chạy trên Docker Desktop Linux containers và chỉ bind dịch vụ ra loopback. Phạm vi gồm DVWA, Nginx, MariaDB, Suricata, Falco, Wazuh manager, collector, normalizer, PostgreSQL, Redis, FastAPI dashboard và các mô hình AI. Pentest chỉ thực hiện trên DVWA local với traffic có marker để truy vết.")
    replace_paragraph(doc.paragraphs[7], "2. Các hạng mục công việc đã thực hiện")
    replace_paragraph(doc.paragraphs[8], "Nghiên cứu cơ sở lý thuyết: Đã tổng hợp kiến trúc IDS/IPS lai, phát hiện dựa trên luật, phát hiện bất thường, tương quan sự kiện, attack graph, vòng đời luật sinh tự động và nguyên tắc response có kiểm soát. Các nguồn log được phân loại theo ba lớp: mạng, ứng dụng và runtime container.")
    replace_paragraph(doc.paragraphs[9], "Triển khai kỹ thuật/Thực hành: Đã xây dựng Compose fragments cho Phase 1 đến Phase 14; tích hợp Suricata EVE JSON, Falco eBPF, Wazuh JSON alerts và Nginx JSONL; triển khai collector, normalizer, Redis Streams và PostgreSQL; hoàn thiện FastAPI/WebSocket dashboard; tích hợp Network AI, Log AI, Graph AI, Fusion, Incident Engine, Rule Generator và Response Engine. Đã bổ sung runbook tiếng Việt, sơ đồ kiến trúc PNG, script kiểm tra interface và script pentest DVWA có login.")
    replace_paragraph(doc.paragraphs[10], "Tiến độ tài liệu: Báo cáo tiến độ đã được cập nhật theo mẫu của bài tập lớn. Tài liệu hướng dẫn kiểm thử, giám sát và pentest đã có. Sơ đồ tổng quan mô hình đã được xuất thành file PNG và liên kết trong README. Phần slide thuyết trình chưa được yêu cầu trong phạm vi file này.")
    replace_paragraph(doc.paragraphs[11], "3. Kết quả bước đầu đạt được")
    replace_paragraph(doc.paragraphs[12], "Sản phẩm hữu hình: Mã nguồn hệ thống AHR-IDPS; cấu hình Docker Compose và sensor; luật Suricata/Falco/Wazuh; backend FastAPI; dashboard React; pipeline thu thập và chuẩn hóa sự kiện; model artifacts; evidence bundle; runbook vận hành; sơ đồ kiến trúc `docs/ahr-idps-model-overview.png`; script kiểm tra interface và script sinh traffic pentest có kiểm soát.")
    replace_paragraph(doc.paragraphs[13], "Chỉ số đánh giá: Bộ static acceptance Phase 1-14 đã được thiết kế và chạy qua các gate kiểm tra cấu hình, security policy, schema, dashboard contract và model contract. Các nhóm unit test của pipeline, backend và AI đã có. Runtime dashboard được kiểm tra qua HTTP API và WebSocket. Một lượt kiểm thử DVWA có login đã tạo được marker trong telemetry, giúp đối chiếu từ request đến event, alert và incident. Response engine mặc định ở chế độ alert_only để tránh phản ứng phá hủy trong lúc nghiên cứu.")
    replace_paragraph(doc.paragraphs[14], "4. Khó khăn và vướng mắc")

    add_body(doc, "Trong quá trình chạy Docker, trạng thái container có thể chuyển sang Exited sau khi Docker Desktop dừng hoặc khởi động lại. Redis có thể cần nhiều bộ nhớ khi nạp lại AOF/RDB và lưu event stream lớn. Vì vậy việc khởi động phải theo thứ tự core → sensor → pipeline → backend → AI; không xóa volume dữ liệu khi xử lý lỗi.")
    add_body(doc, "Dashboard cần xử lý trường hợp Redis hoặc WebSocket chưa sẵn sàng, nếu không một lỗi kết nối có thể làm màn hình hiển thị rỗng. Wazuh hiện được triển khai theo mô hình manager-only để giảm độ phức tạp; indexer và Wazuh Dashboard đầy đủ chưa nằm trong phạm vi nhỏ của bài tập lớn. Việc kiểm soát manager được thực hiện qua log, alert JSON và API manager bind loopback.")
    add_body(doc, "Các mô hình AI phụ thuộc vào dữ liệu có nhãn và chất lượng event. Khi chưa có đủ dữ liệu cho một model, dashboard phải hiển thị trạng thái chưa sẵn sàng thay vì coi là một cảnh báo bảo mật. Đây là điểm cần tiếp tục hoàn thiện trong giai đoạn đánh giá.")

    add_heading(doc, "5. Kế hoạch công việc tiếp theo")
    add_body(doc, "1) Rút gọn Compose và tài liệu vận hành về một lệnh khởi động nhất quán, có kiểm tra Docker daemon và health gate. 2) Hoàn thiện màn hình quản trị Wazuh manager, hiển thị trạng thái daemon, số lượng alert mới, đường dẫn API loopback và lệnh kiểm tra. 3) Chạy ba kịch bản pentest chuẩn hóa trên DVWA, lưu marker, alert, incident và evidence. 4) Đánh giá precision, recall, FPR, latency và khả năng phát hiện unknown attack. 5) Hoàn thiện báo cáo cuối và slide thuyết trình.")

    add_heading(doc, "6. Ba kịch bản pentest dự kiến")
    add_body(doc, "Kịch bản 1 - Baseline hợp lệ: Máy khách đăng nhập DVWA bằng tài khoản lab, truy cập trang chủ, login và instructions. Mục tiêu là đo traffic bình thường, kiểm tra false positive và xác nhận Nginx, Suricata, collector, normalizer và dashboard hoạt động.")
    add_body(doc, "Kịch bản 2 - Web injection: Sau khi đăng nhập, máy khách gửi một tín hiệu SQL injection, reflected XSS, command injection và file inclusion với tốc độ thấp. Mỗi request có scenario ID. Kỳ vọng là Nginx ghi access event, Suricata ghi HTTP/alert, Wazuh nhận alert khi rule khớp, còn Falco ghi process bất thường đối với command injection.")
    add_body(doc, "Kịch bản 3 - Authentication abuse: Máy khách đăng nhập vào DVWA rồi gửi một số lần đăng nhập sai có giới hạn và truy cập một path nhạy cảm trong lab. Mục tiêu là kiểm tra sessionization, phát hiện pattern lặp, tương quan nhiều nguồn và việc hiển thị risk/incident trên dashboard. Không thực hiện DoS, reverse shell, persistence hoặc truy cập ngoài loopback.")

    add_heading(doc, "7. Kết luận tiến độ")
    add_body(doc, "Đề tài đã có nền tảng triển khai từ sensor đến dashboard và có thể dùng để minh họa quy trình giám sát xâm nhập trên DVWA. Phần còn cần hoàn thiện tập trung vào độ ổn định khi Docker restart, quản trị Wazuh manager, xử lý lỗi giao diện và đánh giá định lượng ba kịch bản pentest. Các bước tiếp theo đã được giới hạn để phù hợp với quy mô bài tập lớn và thời gian thực hiện.")

    # Normalize base font for newly added content while preserving the template's layout.
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            if run.text:
                set_font(run, size=run.font.size.pt if run.font.size else 13, bold=run.bold)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
